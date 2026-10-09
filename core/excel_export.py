# core/excel_export.py
"""Exports Excel (.xlsx) :
- du profil projet : les points calculés (nommés, X, Z) et les paramètres de géométrie qui
  les ont produits (export_project_profile) ;
- du profil en long d'un scénario : profils en travers et points durs, par distance
  (export_longitudinal_profile).
Sans dépendance à Qt : la fenêtre principale se contente de demander le chemin et
d'appeler la fonction d'export."""
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.geometry import build_project_cross_section
from core.longitudinal import LongitudinalProfile, longitudinal_table
from core.models import ProjectParameters

# Paramètres de géométrie exportés : (clé, rubrique, libellé, unité), dans l'ordre et avec
# les libellés du formulaire "Profil projet".
_PARAMETERS = (
    ("anchor_x", "Ancrage", "X bord gauche fond", "m"),
    ("anchor_z", "Ancrage", "Z fond du lit", "m NGF"),
    ("bed_width", "Lit trapézoïdal", "Largeur fond", "m"),
    ("bed_depth", "Lit trapézoïdal", "Profondeur", "m"),
    ("bed_side_slope", "Lit trapézoïdal", "Pente bords", "H/V"),
    ("berm_width_left", "Gauche", "Banquette", "m"),
    ("berm_slope_left", "Gauche", "Pente banquette", "m/m"),
    ("bank_slope_left", "Gauche", "Pente berge", "H/V"),
    ("bank_width_left", "Gauche", "Largeur berge", "m"),
    ("floodplain_width_left", "Gauche", "Largeur lit majeur", "m"),
    ("floodplain_slope_left", "Gauche", "Pente lit majeur", "m/m"),
    ("berm_width_right", "Droite", "Banquette", "m"),
    ("berm_slope_right", "Droite", "Pente banquette", "m/m"),
    ("bank_slope_right", "Droite", "Pente berge", "H/V"),
    ("bank_width_right", "Droite", "Largeur berge", "m"),
    ("floodplain_width_right", "Droite", "Largeur lit majeur", "m"),
    ("floodplain_slope_right", "Droite", "Pente lit majeur", "m/m"),
    ("connect_x_left", "Raccord au terrain naturel", "X raccord gauche", "m"),
    ("connect_z_left", "Raccord au terrain naturel", "Z raccord gauche", "m NGF"),
    ("connect_x_right", "Raccord au terrain naturel", "X raccord droit", "m"),
    ("connect_z_right", "Raccord au terrain naturel", "Z raccord droit", "m NGF"),
)

_CONTEXT_LABELS_PROFILE = ("Projet", "Scénario", "Profil")
_CONTEXT_LABELS_DRAFT = ("Zone", "Brouillon")

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4F6D8F")
_NUMBER_FORMAT = "0.000"


def _write_header(sheet, titles: Sequence[str]) -> None:
    sheet.append(list(titles))
    for cell in sheet[sheet.max_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def export_project_profile(path, project_params: Dict[str, Any],
                           context: Optional[Sequence[str]] = None) -> None:
    """Écrit le classeur dans `path`. `project_params` est le dictionnaire du formulaire
    (ou de la base) ; `context` les libellés (projet, scénario, profil) ou ("Draft",
    brouillon) de la sidebar, repris en tête de la feuille des paramètres. Lève
    ValueError si la géométrie est invalide (même règle que le graphique), OSError si
    le fichier ne peut pas être écrit (ex. déjà ouvert dans Excel)."""
    valid_keys = ProjectParameters.__dataclass_fields__.keys()
    params = ProjectParameters(**{k: v for k, v in project_params.items() if k in valid_keys})
    section = build_project_cross_section(params)

    workbook = Workbook()

    # --- Feuille 1 : points du profil projet, de gauche à droite ---
    points_sheet = workbook.active
    points_sheet.title = "Points profil projet"
    _write_header(points_sheet, ("N°", "Point", "X (m)", "Z (m NGF)"))
    for number, point in enumerate(section.points, start=1):
        points_sheet.append([number, point.label, point.x, point.z])
        for cell in points_sheet[points_sheet.max_row][2:]:
            cell.number_format = _NUMBER_FORMAT
    points_sheet.freeze_panes = "A2"

    # --- Feuille 2 : contexte et paramètres de géométrie ---
    params_sheet = workbook.create_sheet("Paramètres")
    if context:
        labels = _CONTEXT_LABELS_DRAFT if len(context) == 2 else _CONTEXT_LABELS_PROFILE
        for label, value in zip(labels, context):
            params_sheet.append([label, value])
            params_sheet.cell(params_sheet.max_row, 1).font = Font(bold=True)
    params_sheet.append(["Exporté le", datetime.now().strftime("%d/%m/%Y %H:%M")])
    params_sheet.cell(params_sheet.max_row, 1).font = Font(bold=True)
    params_sheet.append([])

    _write_header(params_sheet, ("Rubrique", "Paramètre", "Valeur", "Unité"))
    first_param_row = params_sheet.max_row + 1
    for key, group, label, unit in _PARAMETERS:
        value = getattr(params, key)
        params_sheet.append([group, label, "non défini" if value is None else value, unit])
        value_cell = params_sheet.cell(params_sheet.max_row, 3)
        if value is not None:
            value_cell.number_format = _NUMBER_FORMAT if unit not in ("m/m",) else "0.0000"
    params_sheet.freeze_panes = params_sheet.cell(first_param_row, 1)

    for sheet, widths in ((points_sheet, (6, 26, 12, 12)), (params_sheet, (28, 22, 14, 10))):
        for column, width in zip("ABCD", widths):
            sheet.column_dimensions[column].width = width

    workbook.save(Path(path))


# Colonnes du tableau du profil en long : (titre, largeur).
_LONGITUDINAL_COLUMNS = (
    ("Nom", 28), ("Distance existante (m)", 14), ("Distance projet (m)", 14),
    ("Z existant", 14), ("Z projet", 14), ("Z point dur existant", 14), ("Z point dur projet", 14),
)


def export_longitudinal_profile(path, profile: LongitudinalProfile) -> None:
    """Écrit dans `path` le profil en long tel que le graphique l'affiche (cf.
    core.longitudinal.longitudinal_table) : une seule feuille, un seul tableau Excel
    (filtrable), une ligne par profil en travers (ses deux distances) et par point dur (la
    distance et la cote de sa famille), dans l'ordre de longitudinal_table. Une valeur
    absente laisse la case vide. Lève OSError si le fichier ne peut pas être écrit (ex.
    déjà ouvert dans Excel)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Profil en long"
    sheet.append([title for title, _ in _LONGITUDINAL_COLUMNS])
    rows = longitudinal_table(profile)
    for row in rows:
        sheet.append([row.name, row.distance, row.project_distance, row.z_existing, row.z_project,
                      row.z_hard_point, row.z_project_hard_point])
        for cell in sheet[sheet.max_row][1:]:
            cell.number_format = _NUMBER_FORMAT

    # Un tableau Excel couvre au moins une ligne de données, même vide.
    last_column = get_column_letter(len(_LONGITUDINAL_COLUMNS))
    table = Table(displayName="ProfilEnLong", ref=f"A1:{last_column}{max(len(rows), 1) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    for index, (_, width) in enumerate(_LONGITUDINAL_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(Path(path))
