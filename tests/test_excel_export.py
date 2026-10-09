import pytest
from openpyxl import load_workbook

from core.excel_export import export_longitudinal_profile, export_project_profile
from core.longitudinal import build_longitudinal_profile
from core.geometry import build_project_cross_section
from core.models import ProjectParameters


def _rows(sheet):
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_points_sheet_lists_the_named_project_points_left_to_right(tmp_path):
    params = vars(ProjectParameters(floodplain_width_left=2.0, connect_x_right=30.0, connect_z_right=50.0))
    path = tmp_path / "export.xlsx"

    export_project_profile(path, params, ("Rivière", "Variante A", "PK 300"))

    rows = _rows(load_workbook(path)["Points profil projet"])
    assert rows[0] == ["N°", "Point", "X (m)", "Z (m NGF)"]
    assert [r[1] for r in rows[1:]] == [
        "Bout lit majeur G", "Haut de berge G", "Pied de berge G", "Banquette G",
        "Fond du lit G (ancrage)", "Fond du lit D", "Banquette D", "Pied de berge D",
        "Haut de berge D", "Raccord terrain D",
    ]
    expected = build_project_cross_section(ProjectParameters(**params)).points
    assert [(r[2], r[3]) for r in rows[1:]] == [pytest.approx((p.x, p.z)) for p in expected]
    assert [r[0] for r in rows[1:]] == list(range(1, len(expected) + 1))


def test_parameters_sheet_has_context_and_every_geometry_parameter(tmp_path):
    params = vars(ProjectParameters(bed_width=2.5))
    path = tmp_path / "export.xlsx"

    export_project_profile(path, params, ("Rivière", "Variante A", "PK 300"))

    rows = _rows(load_workbook(path)["Paramètres"])
    assert rows[:3] == [["Projet", "Rivière", None, None], ["Scénario", "Variante A", None, None],
                        ["Profil", "PK 300", None, None]]
    assert rows[3][0] == "Exporté le"
    header = rows.index(["Rubrique", "Paramètre", "Valeur", "Unité"])
    by_label = {(r[0], r[1]): (r[2], r[3]) for r in rows[header + 1:]}
    assert by_label[("Lit trapézoïdal", "Largeur fond")] == (2.5, "m")
    assert by_label[("Ancrage", "Z fond du lit")] == (pytest.approx(47.40), "m NGF")
    assert by_label[("Raccord au terrain naturel", "X raccord gauche")] == ("non défini", "m")
    assert len(by_label) == 21


def test_draft_context_is_labelled_as_a_draft(tmp_path):
    path = tmp_path / "export.xlsx"

    export_project_profile(path, vars(ProjectParameters()), ("Draft", "Essai"))

    rows = _rows(load_workbook(path)["Paramètres"])
    assert rows[:2] == [["Zone", "Draft", None, None], ["Brouillon", "Essai", None, None]]


def test_invalid_geometry_raises_and_writes_nothing(tmp_path):
    path = tmp_path / "export.xlsx"
    params = vars(ProjectParameters(connect_x_left=100.0, connect_z_left=50.0))  # côté droit : invalide

    with pytest.raises(ValueError, match="raccord gauche"):
        export_project_profile(path, params)
    assert not path.exists()


# --- Profil en long ---

def _longitudinal_sheet(path):
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Profil en long"]  # une seule feuille
    return workbook["Profil en long"]


def test_longitudinal_export_is_one_table_sorted_by_distance_with_empty_cells(tmp_path):
    hard_points = [{"name": "Pont", "pk": 1000.0, "z": 50.0}, {"name": "Seuil", "pk": 1150.0, "z": 48.0}]
    rows = [(150.0, 47.9, 47.6, "PK 150"), (0.0, 49.8, None, "Amont"), (80.0, None, 48.4, "PK 80")]
    path = tmp_path / "profil_en_long.xlsx"

    export_longitudinal_profile(path, build_longitudinal_profile(rows, hard_points))

    sheet = _longitudinal_sheet(path)
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [
        ["Nom", "Distance (m)", "Z existant", "Z projet", "Z point dur"],
        ["Pont", 0, None, None, 50],
        ["Amont", 0, 49.8, None, None],
        ["PK 80", 80, None, 48.4, None],
        ["Seuil", 150, None, None, 48],
        ["PK 150", 150, 47.9, 47.6, None],
    ]
    assert list(sheet.tables) == ["ProfilEnLong"]  # un seul tableau, sur toutes les lignes
    assert sheet.tables["ProfilEnLong"].ref == "A1:E6"


def test_longitudinal_export_keeps_full_precision(tmp_path):
    """Les valeurs écrites sont celles du graphique, sans arrondi (seul l'affichage l'est)."""
    path = tmp_path / "precision.xlsx"
    export_longitudinal_profile(path, build_longitudinal_profile([(12.3456789, 47.123456789, 46.987654321, "PK 12")]))

    _, row = list(_longitudinal_sheet(path).iter_rows(values_only=True))
    assert row == ("PK 12", 12.3456789, 47.123456789, 46.987654321, None)


def test_longitudinal_export_without_any_row_is_still_a_valid_table(tmp_path):
    path = tmp_path / "vide.xlsx"
    export_longitudinal_profile(path, build_longitudinal_profile([]))

    sheet = _longitudinal_sheet(path)
    assert sheet.tables["ProfilEnLong"].ref == "A1:E2"
