"""Tests de l'arborescence (sidebar), du dialogue de scénario et du routage profil/brouillon
de la fenêtre principale. Qt tourne en mode "offscreen" : aucune fenêtre ne s'affiche, et
les dialogues modaux (saisies, confirmations) sont remplacés par des réponses fixes."""
import json
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
# QtWebEngine (graphique de la fenêtre principale) exige d'être importé avant la
# création de la QApplication.
from PyQt6 import QtWebEngineWidgets  # noqa: F401
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (QApplication, QDialog, QDialogButtonBox, QFileDialog, QInputDialog,
                             QLabel, QMessageBox, QWidget)

from database.db_manager import DatabaseManager, DEFAULT_SCENARIO_NAME
from ui import sidebar as sidebar_module
from ui.dialogs.scenario_dialog import ScenarioDialog
from ui.sidebar import Sidebar, PROJECT, SCENARIO, PROFILE, DRAFT_ROOT, DRAFT


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def db(tmp_path):
    return DatabaseManager(db_path=tmp_path / "test.db")


@pytest.fixture
def sidebar(qapp, db):
    widget = Sidebar(db)
    yield widget
    widget.deleteLater()


def tree_snapshot(sidebar: Sidebar):
    """L'arbre sous forme de listes imbriquées [(type, texte, [enfants])]."""
    def walk(item):
        data = item.data(Qt.ItemDataRole.UserRole)
        children = [walk(item.child(r)) for r in range(item.rowCount())]
        return (data["type"], item.text(), children)

    return [walk(sidebar.model.item(r)) for r in range(sidebar.model.rowCount())]


def click(sidebar: Sidebar, key):
    item = sidebar._find_item(key)
    assert item is not None, key
    index = sidebar.model.indexFromItem(item)
    sidebar.tree_view.setCurrentIndex(index)
    sidebar.on_item_clicked(index)


def record(signal):
    received = []
    signal.connect(lambda *args: received.append(args))
    return received


def answer_inputs(monkeypatch, texts=(), doubles=()):
    """Remplace les saisies QInputDialog par des réponses successives validées."""
    texts, doubles = list(texts), list(doubles)
    monkeypatch.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: (texts.pop(0), True)))
    monkeypatch.setattr(QInputDialog, "getDouble", staticmethod(lambda *a, **k: (doubles.pop(0), True)))


def answer_item_choice(monkeypatch, label: str):
    """Remplace le sélecteur de liste (QInputDialog.getItem) par un choix fixe, validé."""
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(lambda *a, **k: (label, True)))


def forbid_item_choice(monkeypatch):
    """Fait échouer le test si QInputDialog.getItem est appelé (aucune destination ne
    doit être demandée, ex. import d'un fichier projet)."""
    def fail(*a, **k):
        raise AssertionError("QInputDialog.getItem ne devrait pas être appelé ici")
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(fail))


def answer_file_dialogs(monkeypatch, open_path=None, save_path=None):
    """Remplace les sélecteurs de fichier par un chemin fixe, sans ouvrir de fenêtre."""
    if open_path is not None:
        monkeypatch.setattr(QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(open_path), "")))
    if save_path is not None:
        monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(save_path), "")))


# --- Arborescence ---

def test_tree_shows_projects_scenarios_profiles_then_draft_zone(sidebar, db):
    project_id = db.create_project("Rivière")
    scenario_a = db.create_scenario(project_id, "A")
    scenario_b = db.create_scenario(project_id, "B")
    db.create_or_get_profile(scenario_a, "PK 300", 300.0)
    db.create_or_get_profile(scenario_a, "PK 100", 100.0)
    db.create_or_get_profile(scenario_b, "PK 300", 300.0)
    db.create_draft("Essai")

    sidebar.refresh_tree()

    assert tree_snapshot(sidebar) == [
        (PROJECT, "Rivière", [
            (SCENARIO, "A", [(PROFILE, "PK 100", []), (PROFILE, "PK 300", [])]),
            (SCENARIO, "B", [(PROFILE, "PK 300", [])]),
        ]),
        (DRAFT_ROOT, "Draft", [(DRAFT, "Essai", [])]),
    ]


def test_draft_zone_is_present_even_without_projects_or_drafts(sidebar):
    assert tree_snapshot(sidebar) == [(DRAFT_ROOT, "Draft", [])]


def test_new_nodes_are_expanded_and_expansion_survives_refresh(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    sidebar.refresh_tree()
    scenario_index = sidebar.model.indexFromItem(sidebar._find_item((SCENARIO, scenario_id)))
    assert sidebar.tree_view.isExpanded(scenario_index)

    sidebar.tree_view.setExpanded(scenario_index, False)
    sidebar.refresh_tree()

    scenario_index = sidebar.model.indexFromItem(sidebar._find_item((SCENARIO, scenario_id)))
    assert not sidebar.tree_view.isExpanded(scenario_index)


# --- Sélection et signaux ---

def test_clicks_emit_the_signal_of_each_node_type(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    draft_id = db.create_draft("Essai")
    sidebar.refresh_tree()
    signals = {
        "project": record(sidebar.project_selected),
        "scenario": record(sidebar.scenario_selected),
        "profile": record(sidebar.profile_selected),
        "draft": record(sidebar.draft_selected),
        "draft_zone": record(sidebar.draft_zone_selected),
    }

    click(sidebar, (PROJECT, project_id))
    click(sidebar, (SCENARIO, scenario_id))
    click(sidebar, (PROFILE, profile_id))
    assert sidebar.current_context() == ("P", "S", "PK 0")
    click(sidebar, (DRAFT_ROOT, None))
    assert sidebar.current_context() is None
    click(sidebar, (DRAFT, draft_id))
    assert sidebar.current_context() == ("Draft", "Essai")

    assert signals == {
        "project": [(project_id,)], "scenario": [(scenario_id,)], "profile": [(profile_id,)],
        "draft": [(draft_id,)], "draft_zone": [()],
    }


def test_active_path_marks_parents_of_selected_profile(sidebar, db):
    project_id = db.create_project("P")
    scenario_a = db.create_scenario(project_id, "A")
    scenario_b = db.create_scenario(project_id, "B")
    profile_id = db.create_or_get_profile(scenario_a, "PK 0", 0.0)
    sidebar.refresh_tree()

    click(sidebar, (PROFILE, profile_id))

    assert sidebar.is_active({"type": PROJECT, "id": project_id})
    assert sidebar.is_active({"type": SCENARIO, "id": scenario_a})
    assert not sidebar.is_active({"type": SCENARIO, "id": scenario_b})
    # Même id numérique, autre type : pas de confusion.
    assert not sidebar.is_active({"type": DRAFT, "id": profile_id})


def test_deleting_the_open_item_or_a_parent_clears_the_selection(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    sidebar.refresh_tree()
    cleared = record(sidebar.selection_cleared)

    click(sidebar, (PROFILE, profile_id))
    db.delete_scenario(scenario_id)  # le parent du profil ouvert
    sidebar.refresh_tree()

    assert cleared == [()]
    assert sidebar.current_context() is None


def test_deleting_an_unrelated_item_keeps_the_selection(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    kept = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    other = db.create_or_get_profile(scenario_id, "PK 100", 100.0)
    sidebar.refresh_tree()
    cleared = record(sidebar.selection_cleared)

    click(sidebar, (PROFILE, kept))
    db.delete_profile(other)
    sidebar.refresh_tree()

    assert cleared == []
    assert sidebar.current_context() == ("P", "S", "PK 0")
    assert sidebar.tree_view.currentIndex().data(Qt.ItemDataRole.UserRole)["id"] == kept


def test_renaming_the_open_scenario_updates_the_context(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    sidebar.refresh_tree()
    changed = record(sidebar.context_changed)
    click(sidebar, (PROFILE, profile_id))

    db.rename_scenario(scenario_id, "Variante")
    sidebar.refresh_tree()

    assert changed == [()]
    assert sidebar.current_context() == ("P", "Variante", "PK 0")


# --- Actions ---

def test_add_project_creates_it_with_a_default_scenario(sidebar, db, monkeypatch):
    answer_inputs(monkeypatch, texts=["Nouveau"])
    monkeypatch.setattr(sidebar_module.HardPointsDialog, "exec", lambda self: QDialog.DialogCode.Rejected)
    selected = record(sidebar.scenario_selected)

    sidebar.add_project()

    project = db.get_all_projects()[0]
    assert [s["name"] for s in project["scenarios"]] == [DEFAULT_SCENARIO_NAME]
    assert selected == [(project["scenarios"][0]["id"],)]


def test_add_profile_goes_to_the_selected_scenario(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    scenario_a = db.create_scenario(project_id, "A")
    scenario_b = db.create_scenario(project_id, "B")
    sidebar.refresh_tree()
    click(sidebar, (SCENARIO, scenario_b))
    answer_inputs(monkeypatch, texts=["PK 300"], doubles=[300.0])

    sidebar.add_profile()

    scenarios = db.get_all_projects()[0]["scenarios"]
    assert scenarios[0]["profiles"] == []
    assert [p["name"] for p in scenarios[1]["profiles"]] == ["PK 300"]


def test_add_profile_on_project_with_several_scenarios_asks_to_choose(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    db.create_scenario(project_id, "A")
    db.create_scenario(project_id, "B")
    sidebar.refresh_tree()
    click(sidebar, (PROJECT, project_id))
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a[2])))

    sidebar.add_profile()

    assert warnings and "scénario" in warnings[0]
    assert all(s["profiles"] == [] for s in db.get_all_projects()[0]["scenarios"])


def test_add_profile_in_draft_zone_creates_a_draft_not_a_profile(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    db.create_scenario(project_id, "S")
    sidebar.refresh_tree()
    click(sidebar, (DRAFT_ROOT, None))
    answer_inputs(monkeypatch, texts=["Essai berge"])
    opened = record(sidebar.draft_selected)

    sidebar.add_profile()

    drafts = db.get_all_drafts()
    assert [d["name"] for d in drafts] == ["Essai berge"]
    assert opened == [(drafts[0]["id"],)]
    assert db.get_all_projects()[0]["scenarios"][0]["profiles"] == []


class _FakeScenarioDialog:
    """Remplace ScenarioDialog : valide avec les valeurs choisies par le test, et garde
    trace des arguments reçus (scénarios proposés, source pré-choisie)."""
    answer = None
    received = None

    def __init__(self, parent=None, scenarios=None, default_name="", source_scenario_id=None, title=""):
        type(self).received = {"scenarios": scenarios, "default_name": default_name,
                               "source_scenario_id": source_scenario_id}

    def exec(self):
        return QDialog.DialogCode.Accepted

    def get_values(self):
        return type(self).answer


def test_create_scenario_keeping_profiles_of_an_existing_one(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    source = db.create_scenario(project_id, DEFAULT_SCENARIO_NAME)
    db.create_or_get_profile(source, "PK 0", 0.0)
    db.create_or_get_profile(source, "PK 300", 300.0)
    sidebar.refresh_tree()
    click(sidebar, (SCENARIO, source))
    monkeypatch.setattr(sidebar_module, "ScenarioDialog", _FakeScenarioDialog)
    _FakeScenarioDialog.answer = ("Variante", source)

    sidebar.add_scenario()

    assert _FakeScenarioDialog.received["scenarios"] == [{"id": source, "name": DEFAULT_SCENARIO_NAME}]
    assert _FakeScenarioDialog.received["default_name"] == "Scénario 2"
    scenarios = db.get_all_projects()[0]["scenarios"]
    assert [(s["name"], [p["name"] for p in s["profiles"]]) for s in scenarios] == [
        (DEFAULT_SCENARIO_NAME, ["PK 0", "PK 300"]),
        ("Variante", ["PK 0", "PK 300"]),
    ]


def test_duplicate_scenario_preselects_it_as_source(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    source = db.create_scenario(project_id, "A")
    db.create_or_get_profile(source, "PK 0", 0.0)
    sidebar.refresh_tree()
    monkeypatch.setattr(sidebar_module, "ScenarioDialog", _FakeScenarioDialog)
    _FakeScenarioDialog.answer = ("A - copie", source)

    sidebar.duplicate_scenario({"type": SCENARIO, "id": source, "project_id": project_id}, "A")

    assert _FakeScenarioDialog.received["source_scenario_id"] == source
    assert _FakeScenarioDialog.received["default_name"] == "A - copie"
    assert [s["name"] for s in db.get_scenarios(project_id)] == ["A", "A - copie"]


def test_add_scenario_without_project_selected_warns(sidebar, db, monkeypatch):
    db.create_draft("Essai")
    sidebar.refresh_tree()
    click(sidebar, (DRAFT_ROOT, None))
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a[2])))

    sidebar.add_scenario()

    assert warnings == ["Sélectionnez d'abord un projet."]


# --- Dialogue de scénario ---

def test_scenario_dialog_without_existing_scenario_cannot_keep_profiles(qapp):
    dialog = ScenarioDialog(scenarios=[], default_name="Scénario 1")

    assert not dialog.keep_profiles_check.isEnabled()
    assert dialog.get_values() == ("Scénario 1", None)


def test_scenario_dialog_returns_source_only_when_option_checked(qapp):
    scenarios = [{"id": 4, "name": "A"}, {"id": 7, "name": "B"}]
    dialog = ScenarioDialog(scenarios=scenarios, default_name="  Nouveau  ")
    assert not dialog.source_combo.isEnabled()
    assert dialog.get_values() == ("Nouveau", None)

    dialog.keep_profiles_check.setChecked(True)
    dialog.source_combo.setCurrentIndex(1)

    assert dialog.source_combo.isEnabled()
    assert dialog.get_values() == ("Nouveau", 7)


def test_scenario_dialog_preselected_source_and_empty_name(qapp):
    dialog = ScenarioDialog(scenarios=[{"id": 4, "name": "A"}, {"id": 7, "name": "B"}],
                            default_name="B - copie", source_scenario_id=7)
    assert dialog.get_values() == ("B - copie", 7)

    dialog.name_edit.setText("   ")
    assert not dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled()


# --- Fenêtre principale : un brouillon s'enregistre dans les brouillons, jamais ailleurs ---

class _FakePlotView(QWidget):
    """Remplace PlotView : son QWebEngineView (Chromium) ne démarre pas sans affichage et
    fait tomber le processus en mode offscreen. On ne garde que ce que MainWindow utilise."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.lbl_title = QLabel()
        self.figures = []
        self.view_keys = []

    def update_plot(self, fig, error_message=None, view_key=None):
        self.figures.append((fig, error_message))
        self.view_keys.append(view_key)


@pytest.fixture
def main_window(qapp, tmp_path, monkeypatch):
    from ui import main_window as main_window_module
    db_path = tmp_path / "window.db"
    # Jamais la vraie base (data/hydrotopo.db) : la fenêtre ouvre une base de test.
    monkeypatch.setattr(main_window_module, "DatabaseManager", lambda: DatabaseManager(db_path=db_path))
    monkeypatch.setattr(main_window_module, "PlotView", _FakePlotView)
    window = main_window_module.MainWindow()
    yield window
    window.deleteLater()


def test_editing_a_draft_never_touches_the_profile_with_the_same_id(main_window):
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    profile_state = ([{"X (m)": 0.0, "Z (m NGF)": 1.0}, {"X (m)": 1.0, "Z (m NGF)": 0.5}], {"anchor_z": 0.2})
    db.save_profile_state(profile_id, *profile_state)
    draft_id = db.create_draft("Essai")
    assert draft_id == profile_id
    main_window.sidebar.refresh_tree()

    click(main_window.sidebar, (DRAFT, draft_id))
    assert main_window._current_target == ("draft", draft_id)
    assert main_window.lbl_context.text().count("&rsaquo;") == 1  # Draft › Essai
    draft_points = [{"X (m)": 0.0, "Z (m NGF)": 9.0}, {"X (m)": 4.0, "Z (m NGF)": 8.0}]
    main_window.form_existing.set_data(draft_points)
    main_window.save_and_update_plot()

    assert db.load_draft_state(draft_id)[0] == draft_points
    assert db.load_profile_state(profile_id) == profile_state

    # Et l'inverse : le profil s'ouvre avec ses propres données.
    click(main_window.sidebar, (PROFILE, profile_id))
    assert main_window._current_target == ("profile", profile_id)
    assert main_window.form_existing.get_data() == profile_state[0]
    assert main_window.lbl_context.text().count("&rsaquo;") == 2  # P › S › PK 0


def test_each_profile_draft_and_tab_gets_its_own_view_key(main_window):
    """Le graphique mémorise le zoom par clé de vue : une clé par (profil ou brouillon,
    onglet), jamais partagée entre un profil et un brouillon de même id, et aucune clé pour
    le profil en long (qui ne mémorise rien)."""
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    draft_id = db.create_draft("Essai")
    assert draft_id == profile_id
    points = [{"X (m)": 0.0, "Z (m NGF)": 1.0}, {"X (m)": 1.0, "Z (m NGF)": 0.5}]
    db.save_profile_state(profile_id, points, {})
    db.save_draft_state(draft_id, points, {})
    main_window.sidebar.refresh_tree()
    fake = main_window.plot_view

    click(main_window.sidebar, (PROFILE, profile_id))
    assert fake.view_keys[-1] == f"profile:{profile_id}:existing"

    main_window.tabs.setCurrentIndex(1)  # Profil projet
    assert fake.view_keys[-1] == f"profile:{profile_id}:project"
    main_window.tabs.setCurrentIndex(2)  # Hydraulique
    assert fake.view_keys[-1] == f"profile:{profile_id}:hydraulics"

    click(main_window.sidebar, (SCENARIO, scenario_id))
    assert fake.view_keys[-1] is None

    main_window.tabs.setCurrentIndex(0)
    click(main_window.sidebar, (DRAFT, draft_id))
    assert fake.view_keys[-1] == f"draft:{draft_id}:existing"

    # Retour au même profil : exactement la même clé qu'avant le passage par le profil en long.
    click(main_window.sidebar, (PROFILE, profile_id))
    assert fake.view_keys[-1] == f"profile:{profile_id}:existing"


def test_export_excel_button_writes_the_open_profile(main_window, monkeypatch, tmp_path):
    from openpyxl import load_workbook
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    db.save_profile_state(profile_id, [], {"bed_width": 3.25})
    main_window.sidebar.refresh_tree()
    click(main_window.sidebar, (PROFILE, profile_id))
    target = tmp_path / "sortie.xlsx"
    proposed = []
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda _parent, _title, name, _filter: (proposed.append(name), (str(target), ""))[1]))

    main_window.form_project.btn_export_excel.click()

    assert proposed == ["PK 0 - profil projet.xlsx"]
    rows = list(load_workbook(target)["Paramètres"].iter_rows(values_only=True))
    assert rows[2][:2] == ("Profil", "PK 0")
    assert ("Lit trapézoïdal", "Largeur fond", 3.25, "m") in rows


def test_export_image_name_follows_what_is_displayed(main_window):
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 120", 120.0)
    db.save_profile_state(profile_id, [{"X (m)": 0.0, "Z (m NGF)": 1.0}, {"X (m)": 1.0, "Z (m NGF)": 0.5}], {})
    main_window.sidebar.refresh_tree()

    click(main_window.sidebar, (PROFILE, profile_id))
    assert main_window.plot_view.export_name == "PK 120 - Profil existant"
    main_window.tabs.setCurrentIndex(2)
    assert main_window.plot_view.export_name == "PK 120 - Hydraulique"
    click(main_window.sidebar, (SCENARIO, scenario_id))
    assert main_window.plot_view.export_name == "Profil en long"


POINTS_2 = [{"X (m)": 0.0, "Z (m NGF)": 1.0}, {"X (m)": 1.0, "Z (m NGF)": 0.5}]


def _overlay_boxes(window):
    return window.form_project.chk_overlay.isChecked(), window.form_hydraulics.chk_overlay.isChecked()


@pytest.mark.parametrize("project_box,hydraulics_box", [(True, False), (False, True)])
def test_each_background_profile_box_keeps_its_own_value_after_a_round_trip(main_window, project_box, hydraulics_box):
    """Cocher "Afficher le profil ... en fond" dans un onglet, aller voir le profil en long
    puis revenir : chaque case retrouve SA valeur (avant, celle de l'onglet Hydraulique
    écrasait celle du Profil projet à l'enregistrement)."""
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    db.save_profile_state(profile_id, POINTS_2, {})
    main_window.sidebar.refresh_tree()
    click(main_window.sidebar, (PROFILE, profile_id))

    main_window.form_project.chk_overlay.setChecked(project_box)
    main_window.form_hydraulics.chk_overlay.setChecked(hydraulics_box)
    click(main_window.sidebar, (SCENARIO, scenario_id))
    click(main_window.sidebar, (PROFILE, profile_id))

    assert _overlay_boxes(main_window) == (project_box, hydraulics_box)
    saved = db.load_profile_state(profile_id)[1]
    assert (saved["show_overlay_project"], saved["show_overlay_hydraulics"]) == (project_box, hydraulics_box)


def test_background_profile_boxes_are_kept_for_drafts_too(main_window):
    db = main_window.db_manager
    draft_id = db.create_draft("Essai")
    other_id = db.create_draft("Autre")
    db.save_draft_state(draft_id, POINTS_2, {})
    main_window.sidebar.refresh_tree()
    click(main_window.sidebar, (DRAFT, draft_id))

    main_window.form_project.chk_overlay.setChecked(True)
    main_window.form_hydraulics.chk_overlay.setChecked(False)
    click(main_window.sidebar, (DRAFT, other_id))
    click(main_window.sidebar, (DRAFT, draft_id))

    assert _overlay_boxes(main_window) == (True, False)


@pytest.mark.parametrize("legacy_value", [True, False])
def test_profile_saved_with_the_old_shared_key_initialises_both_boxes(main_window, legacy_value):
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    db.save_profile_state(profile_id, POINTS_2, {"show_overlay": legacy_value})
    main_window.sidebar.refresh_tree()

    click(main_window.sidebar, (PROFILE, profile_id))

    assert _overlay_boxes(main_window) == (legacy_value, legacy_value)


def test_project_and_hydraulics_forms_share_no_saved_key(main_window):
    """Les deux formulaires sont fusionnés dans un seul enregistrement : une clé commune
    ferait écraser la valeur de l'un par celle de l'autre."""
    shared = set(main_window.form_project.get_data()) & set(main_window.form_hydraulics.get_data())
    assert shared == set()


def test_scenario_click_shows_longitudinal_and_stops_editing(main_window):
    db = main_window.db_manager
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    main_window.sidebar.refresh_tree()

    click(main_window.sidebar, (PROFILE, profile_id))
    click(main_window.sidebar, (SCENARIO, scenario_id))

    assert main_window._current_target is None
    assert main_window.plot_view.lbl_title.text() == "Profil en long du scénario"


def test_deleting_the_open_draft_returns_to_the_welcome_page(main_window):
    db = main_window.db_manager
    draft_id = db.create_draft("Essai")
    main_window.sidebar.refresh_tree()
    click(main_window.sidebar, (DRAFT, draft_id))

    db.delete_draft(draft_id)
    main_window.sidebar.refresh_tree()

    assert main_window._current_target is None
    assert main_window.forms_stack.currentIndex() == 0


# --- Export / Import (JSON) ---

def test_export_profile_item_writes_a_profile_file(sidebar, db, tmp_path, monkeypatch):
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    db.save_profile_state(profile_id, [{"X (m)": 0.0, "Z (m NGF)": 1.0}], {"anchor_z": 0.5})
    sidebar.refresh_tree()
    path = tmp_path / "export.json"
    answer_file_dialogs(monkeypatch, save_path=path)

    sidebar.export_profile_item({"type": PROFILE, "id": profile_id}, "PK 0")

    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["type"] == "profile" and data["name"] == "PK 0" and data["distance"] == 0.0


def test_export_draft_item_writes_a_profile_file_without_distance(sidebar, db, tmp_path, monkeypatch):
    draft_id = db.create_draft("Essai")
    sidebar.refresh_tree()
    path = tmp_path / "export.json"
    answer_file_dialogs(monkeypatch, save_path=path)

    sidebar.export_draft_item({"type": DRAFT, "id": draft_id}, "Essai")

    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["type"] == "profile" and data["distance"] is None


def test_export_scenario_and_project_items_write_matching_types(sidebar, db, tmp_path, monkeypatch):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    sidebar.refresh_tree()

    scenario_path = tmp_path / "scenario.json"
    answer_file_dialogs(monkeypatch, save_path=scenario_path)
    sidebar.export_scenario_item({"type": SCENARIO, "id": scenario_id}, "S")
    assert json.loads(scenario_path.read_text(encoding="utf-8"))["type"] == "scenario"

    project_path = tmp_path / "project.json"
    answer_file_dialogs(monkeypatch, save_path=project_path)
    sidebar.export_project_item({"type": PROJECT, "id": project_id}, "P")
    assert json.loads(project_path.read_text(encoding="utf-8"))["type"] == "project"


def test_import_file_profile_routes_to_the_chosen_scenario(sidebar, db, tmp_path, monkeypatch):
    other_project = db.create_project("Autre projet")
    other_scenario = db.create_scenario(other_project, "Source")
    profile_id = db.create_or_get_profile(other_scenario, "PK 0", 0.0)
    db.save_profile_state(profile_id, [{"X (m)": 1.0, "Z (m NGF)": 2.0}], {"anchor_z": 1.5})
    path = tmp_path / "profile.json"
    db.export_profile_to_file(profile_id, path)

    target_project = db.create_project("Cible")
    db.create_scenario(target_project, "Destination")
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    answer_item_choice(monkeypatch, "Cible › Destination")
    opened = record(sidebar.profile_selected)

    sidebar.import_file()

    profiles = next(p for p in db.get_all_projects()
                     if p["name"] == "Cible")["scenarios"][0]["profiles"]
    assert [p["name"] for p in profiles] == ["PK 0"]
    assert len(opened) == 1 and opened[0][0] == profiles[0]["id"]


def test_import_file_profile_routes_to_draft_zone(sidebar, db, tmp_path, monkeypatch):
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    path = tmp_path / "profile.json"
    db.export_profile_to_file(profile_id, path)
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    answer_item_choice(monkeypatch, "Draft (brouillons)")

    sidebar.import_file()

    assert [d["name"] for d in db.get_all_drafts()] == ["PK 0"]


def test_import_file_scenario_routes_to_the_chosen_project(sidebar, db, tmp_path, monkeypatch):
    source_project = db.create_project("Source")
    source_scenario = db.create_scenario(source_project, "Variante")
    db.create_or_get_profile(source_scenario, "PK 0", 0.0)
    path = tmp_path / "scenario.json"
    db.export_scenario_to_file(source_scenario, path)

    db.create_project("Cible")
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    answer_item_choice(monkeypatch, "Cible")
    opened = record(sidebar.scenario_selected)

    sidebar.import_file()

    scenarios = next(p for p in db.get_all_projects() if p["name"] == "Cible")["scenarios"]
    assert [s["name"] for s in scenarios] == ["Variante"]
    assert opened and opened[0][0] == scenarios[0]["id"]


def test_import_file_project_creates_a_new_project_without_prompting(sidebar, db, tmp_path, monkeypatch):
    project_id = db.create_project("Rivière")
    scenario_id = db.create_scenario(project_id, "S")
    db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    path = tmp_path / "project.json"
    db.export_project_to_file(project_id, path)
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    forbid_item_choice(monkeypatch)  # un projet ne demande jamais de destination
    opened = record(sidebar.project_selected)

    sidebar.import_file()

    names = sorted(p["name"] for p in db.get_all_projects())
    assert names == ["Rivière", "Rivière - importé"]
    assert opened


def test_import_profile_into_scenario_context_menu_skips_destination_prompt(sidebar, db, tmp_path, monkeypatch, info_messages):
    draft_id = db.create_draft("Essai")
    db.save_draft_state(draft_id, [{"X (m)": 3.0, "Z (m NGF)": 4.0}], {"anchor_z": 3.5})
    path = tmp_path / "profile.json"
    db.export_draft_to_file(draft_id, path)

    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    forbid_item_choice(monkeypatch)  # le scénario cible est déjà connu (menu contextuel)

    sidebar.import_profile_into_scenario_item({"type": SCENARIO, "id": scenario_id})

    profiles = db.get_all_projects()[0]["scenarios"][0]["profiles"]
    assert [p["name"] for p in profiles] == ["Essai"]
    # Projet sans points durs : le profil importé (pente calculée par défaut) passe en
    # pente imposée, et un message le signale.
    assert db.load_profile_state(profiles[0]["id"]) == (
        [{"X (m)": 3.0, "Z (m NGF)": 4.0}], {"anchor_z": 3.5, "slope_mode": "imposed"})
    assert len(info_messages) == 1 and "passé(s) en pente imposée" in info_messages[0][1]


def test_import_scenario_into_project_context_menu_skips_destination_prompt(sidebar, db, tmp_path, monkeypatch):
    source_scenario = db.create_scenario(db.create_project("Source"), "Variante")
    db.create_or_get_profile(source_scenario, "PK 0", 0.0)
    path = tmp_path / "scenario.json"
    db.export_scenario_to_file(source_scenario, path)

    target_project = db.create_project("Cible")
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    forbid_item_choice(monkeypatch)

    sidebar.import_scenario_into_project_item({"type": PROJECT, "id": target_project})

    scenarios = db.get_scenarios(target_project)
    assert [s["name"] for s in scenarios] == ["Variante"]


def test_import_profile_into_drafts_context_menu(sidebar, db, tmp_path, monkeypatch):
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    path = tmp_path / "profile.json"
    db.export_profile_to_file(profile_id, path)
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    forbid_item_choice(monkeypatch)

    sidebar.import_profile_into_drafts_item()

    assert [d["name"] for d in db.get_all_drafts()] == ["PK 0"]


def test_import_wrong_file_type_in_context_menu_warns_and_imports_nothing(sidebar, db, tmp_path, monkeypatch):
    """Importer un fichier "projet" via le menu contextuel d'un scénario (qui n'attend
    qu'un fichier "profil") : avertit, ne crée rien."""
    project_id = db.create_project("Source")
    path = tmp_path / "project.json"
    db.export_project_to_file(project_id, path)
    scenario_id = db.create_scenario(db.create_project("P"), "S")
    sidebar.refresh_tree()
    answer_file_dialogs(monkeypatch, open_path=path)
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a[2])))

    sidebar.import_profile_into_scenario_item({"type": SCENARIO, "id": scenario_id})

    assert warnings and "profil" in warnings[0]
    assert db.get_all_projects()[0]["scenarios"][0]["profiles"] == []


# --- Copie interne Draft → Scénario ---

def test_copy_draft_to_scenario_action_creates_an_independent_copy(sidebar, db, monkeypatch):
    draft_id = db.create_draft("Essai berge")
    db.save_draft_state(draft_id, [{"X (m)": 0.0, "Z (m NGF)": 5.0}], {"anchor_z": 4.5})
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    sidebar.refresh_tree()
    answer_item_choice(monkeypatch, "P › S")
    opened = record(sidebar.profile_selected)

    sidebar.copy_draft_to_scenario({"type": DRAFT, "id": draft_id}, "Essai berge")

    profiles = db.get_all_projects()[0]["scenarios"][0]["profiles"]
    assert [p["name"] for p in profiles] == ["Essai berge"]
    assert opened and opened[0][0] == profiles[0]["id"]
    # Le brouillon source reste inchangé dans la zone Draft.
    assert db.get_all_drafts() == [{"id": draft_id, "name": "Essai berge"}]
    assert db.load_draft_state(draft_id) == ([{"X (m)": 0.0, "Z (m NGF)": 5.0}], {"anchor_z": 4.5})


def test_copy_draft_to_scenario_without_any_scenario_warns(sidebar, db, monkeypatch):
    draft_id = db.create_draft("Essai")
    sidebar.refresh_tree()
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a[2])))

    sidebar.copy_draft_to_scenario({"type": DRAFT, "id": draft_id}, "Essai")

    assert warnings == ["Aucun scénario : créez d'abord un projet et un scénario."]


# --- Archives ---

def test_archive_action_removes_the_project_from_the_tree_and_counts_it(sidebar, db):
    project_id = db.create_project("P")
    db.create_scenario(project_id, "S")
    db.create_project("Autre")
    sidebar.refresh_tree()
    assert sidebar.btn_archives.text() == "Archives"

    sidebar.archive_project_item({"type": PROJECT, "id": project_id})

    assert [name for _, name, _ in tree_snapshot(sidebar)] == ["Autre", "Draft"]
    assert sidebar.btn_archives.text() == "Archives (1)"
    assert [p["id"] for p in db.get_archived_projects()] == [project_id]


def test_archiving_the_open_project_clears_the_selection(sidebar, db):
    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    sidebar.refresh_tree()
    click(sidebar, (PROFILE, profile_id))
    cleared = record(sidebar.selection_cleared)

    sidebar.archive_project_item({"type": PROJECT, "id": project_id})

    assert cleared == [()]
    assert sidebar.current_context() is None


def test_archived_projects_are_not_offered_as_import_or_copy_destinations(sidebar, db):
    kept = db.create_project("Actif")
    db.create_scenario(kept, "S")
    archived = db.create_project("Rangé")
    db.create_scenario(archived, "S")
    db.set_project_archived(archived, True)

    assert sidebar._project_choices() == [("Actif", kept)]
    assert [label for label, _ in sidebar._scenario_choices()] == ["Actif › S"]


def test_archives_dialog_restores_and_the_sidebar_selects_the_project(sidebar, db, monkeypatch):
    from ui.dialogs.archives_dialog import ArchivesDialog
    project_id = db.create_project("P")
    db.create_scenario(project_id, "S")
    db.set_project_archived(project_id, True)
    sidebar.refresh_tree()
    opened = record(sidebar.project_selected)

    def fake_exec(self):
        assert self.list_widget.count() == 1
        self.restore_selected()
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ArchivesDialog, "exec", fake_exec)

    sidebar.open_archives()

    assert [name for _, name, _ in tree_snapshot(sidebar)] == ["P", "Draft"]
    assert sidebar.btn_archives.text() == "Archives"
    assert opened == [(project_id,)]


def test_archives_dialog_lists_contents_and_disables_actions_when_empty(qapp, db):
    from ui.dialogs.archives_dialog import ArchivesDialog
    empty = ArchivesDialog(db)
    assert empty.list_widget.count() == 0
    assert not empty.btn_restore.isEnabled() and not empty.btn_delete.isEnabled()

    project_id = db.create_project("P")
    scenario_id = db.create_scenario(project_id, "S")
    db.create_or_get_profile(scenario_id, "PK 0", 0.0)
    db.set_project_archived(project_id, True)
    dialog = ArchivesDialog(db)

    assert dialog.list_widget.item(0).text() == "P  —  1 scénario, 1 profil"
    assert dialog.btn_restore.isEnabled() and dialog.btn_delete.isEnabled()


def test_archives_dialog_deletes_only_after_confirmation(qapp, db, monkeypatch):
    from ui.dialogs.archives_dialog import ArchivesDialog
    project_id = db.create_project("P")
    db.set_project_archived(project_id, True)
    dialog = ArchivesDialog(db)

    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.No))
    dialog.delete_selected()
    assert [p["id"] for p in db.get_archived_projects()] == [project_id]

    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes))
    dialog.delete_selected()
    assert db.get_archived_projects() == []
    assert dialog.list_widget.count() == 0


# --- Points durs multiples et pente hydraulique calculée ---

HP_POINTS = [
    {"name": "A", "pk": 1000.0, "z": 50.0},
    {"name": "B", "pk": 1200.0, "z": 46.0},
    {"name": "C", "pk": 1400.0, "z": 45.0},
]


def _accept_hard_points_dialog(monkeypatch, edit):
    """Remplace l'affichage du dialogue des points durs : `edit(dialog)` modifie le
    tableau comme le ferait l'utilisateur, puis le dialogue est validé (s'il n'y a pas
    d'erreur, comme avec le vrai bouton OK)."""
    from PyQt6.QtWidgets import QDialogButtonBox
    from ui.dialogs.hard_points_dialog import HardPointsDialog

    def fake_exec(dialog):
        edit(dialog)
        ok = dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled()
        return QDialog.DialogCode.Accepted if ok else QDialog.DialogCode.Rejected
    monkeypatch.setattr(HardPointsDialog, "exec", fake_exec)


def test_editing_hard_points_recomputes_slopes_and_reports_them(sidebar, db, monkeypatch, info_messages):
    from ui.dialogs.hard_points_dialog import COL_Z
    project_id = db.create_project("P")
    db.set_hard_points(project_id, HP_POINTS)
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK 1100", 100.0)
    sidebar.refresh_tree()
    changed = record(sidebar.project_data_changed)
    _accept_hard_points_dialog(monkeypatch, lambda d: d.table.item(0, COL_Z).setText("52"))

    sidebar.edit_hard_points({"type": PROJECT, "id": project_id})

    assert db.get_hard_points(project_id)[0]["z"] == 52.0
    assert db.load_profile_state(profile_id)[1]["slope"] == pytest.approx(0.03)
    assert info_messages and "1 profil(s) mis à jour : S › PK 1100" in info_messages[-1][1]
    assert changed == [(project_id,)]


def test_hard_points_that_would_leave_a_profile_outside_cannot_be_validated(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    db.set_hard_points(project_id, HP_POINTS)
    scenario_id = db.create_scenario(project_id, "S")
    db.create_or_get_profile(scenario_id, "PK 1350", 350.0)
    sidebar.refresh_tree()
    seen = {}

    def remove_last_point(dialog):
        dialog.table.selectRow(2)
        dialog.remove_selected_point()
        seen["errors"] = dialog.lbl_errors.text()
    _accept_hard_points_dialog(monkeypatch, remove_last_point)

    sidebar.edit_hard_points({"type": PROJECT, "id": project_id})

    assert "« S › PK 1350 » (PK 1350)" in seen["errors"]
    assert len(db.get_hard_points(project_id)) == 3  # rien d'enregistré


def test_distance_prompt_is_bounded_by_the_hard_points(sidebar, db, monkeypatch):
    project_id = db.create_project("P")
    db.set_hard_points(project_id, HP_POINTS)
    scenario_id = db.create_scenario(project_id, "S")
    sidebar.refresh_tree()
    asked = {}

    def fake_get_double(parent, title, label, value, low, high, decimals):
        asked.update(label=label, low=low, high=high)
        return 250.0, True
    monkeypatch.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: ("PK 1250", True)))
    monkeypatch.setattr(QInputDialog, "getDouble", staticmethod(fake_get_double))

    sidebar.add_profile_to_scenario(scenario_id, project_id)

    assert (asked["low"], asked["high"]) == (0.0, 400.0)
    assert "au premier point dur « A »" in asked["label"] and "entre 0 et 400" in asked["label"]
    profile = db.get_all_projects()[0]["scenarios"][0]["profiles"][0]
    assert db.load_profile_state(profile["id"])[1] == {"slope_mode": "computed", "slope": pytest.approx(0.005)}


def _open_profile(window, distance=100.0, points=HP_POINTS):
    db = window.db_manager
    project_id = db.create_project("P")
    db.set_hard_points(project_id, points)
    scenario_id = db.create_scenario(project_id, "S")
    profile_id = db.create_or_get_profile(scenario_id, "PK", distance)
    window.sidebar.refresh_tree()
    click(window.sidebar, (PROFILE, profile_id))
    return project_id, profile_id


def test_hydraulics_tab_shows_the_computed_slope_greyed_with_its_segment(main_window):
    _open_profile(main_window, distance=250.0)
    form = main_window.form_hydraulics

    assert not form.chk_impose_slope.isChecked() and form.chk_impose_slope.isEnabled()
    assert not form.inputs["slope"].isEnabled()
    assert form.inputs["slope"].value() == pytest.approx(0.005)
    assert form.lbl_slope_info.text() == "Pente calculée sur le tronçon « B » → « C » : 0.0050 m/m."


def test_imposing_a_slope_then_unchecking_returns_to_the_computed_value(main_window):
    _, profile_id = _open_profile(main_window, distance=100.0)
    form = main_window.form_hydraulics
    db = main_window.db_manager

    form.chk_impose_slope.setChecked(True)
    form.inputs["slope"].setValue(0.0123)
    assert form.inputs["slope"].isEnabled()
    assert db.load_profile_state(profile_id)[1]["slope_mode"] == "imposed"
    assert db.load_profile_state(profile_id)[1]["slope"] == pytest.approx(0.0123)
    assert "Pente calculée : 0.0200 m/m" in form.lbl_slope_info.text()

    form.chk_impose_slope.setChecked(False)
    assert form.inputs["slope"].value() == pytest.approx(0.02) and not form.inputs["slope"].isEnabled()
    saved = db.load_profile_state(profile_id)[1]
    assert (saved["slope_mode"], saved["slope"]) == ("computed", pytest.approx(0.02))


def test_imposed_slope_is_kept_when_the_profile_is_reopened(main_window):
    _, profile_id = _open_profile(main_window)
    form = main_window.form_hydraulics
    form.chk_impose_slope.setChecked(True)
    form.inputs["slope"].setValue(0.0123)

    click(main_window.sidebar, (SCENARIO, main_window.db_manager.get_scenarios(1)[0]["id"]))
    click(main_window.sidebar, (PROFILE, profile_id))

    assert form.chk_impose_slope.isChecked() and form.inputs["slope"].value() == pytest.approx(0.0123)


def test_without_possible_computation_the_slope_is_imposed_and_explained(main_window):
    _open_profile(main_window, points=[{"name": "Seul", "pk": 0.0, "z": 50.0}])
    form = main_window.form_hydraulics

    assert form.chk_impose_slope.isChecked() and not form.chk_impose_slope.isEnabled()
    assert form.inputs["slope"].isEnabled()
    assert "calcul impossible" in form.lbl_slope_info.text()
    assert "au moins deux points durs" in form.lbl_slope_info.text()


def test_drafts_always_have_an_imposed_slope(main_window):
    db = main_window.db_manager
    draft_id = db.create_draft("Essai")
    main_window.sidebar.refresh_tree()

    click(main_window.sidebar, (DRAFT, draft_id))

    form = main_window.form_hydraulics
    assert form.chk_impose_slope.isChecked() and not form.chk_impose_slope.isEnabled()
    assert "Brouillon" in form.lbl_slope_info.text()
    assert db.load_draft_state(draft_id)[1].get("slope_mode") in (None, "imposed")


def test_open_profile_is_refreshed_after_its_hard_points_change(main_window, monkeypatch):
    from ui.dialogs.hard_points_dialog import COL_Z
    project_id, _ = _open_profile(main_window, distance=100.0)
    _accept_hard_points_dialog(monkeypatch, lambda d: d.table.item(0, COL_Z).setText("52"))

    main_window.sidebar.edit_hard_points({"type": PROJECT, "id": project_id})

    assert main_window.form_hydraulics.inputs["slope"].value() == pytest.approx(0.03)


def test_startup_message_reports_the_migration(main_window, info_messages):
    main_window.db_manager.startup_slope_report.updated = ["S › PK 100"]

    main_window.show_startup_messages()

    assert info_messages[-1][0] == "Mise à jour des pentes hydrauliques"
    assert "1 profil(s) mis à jour : S › PK 100" in info_messages[-1][1]
